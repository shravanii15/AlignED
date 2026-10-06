"""Tests for resume text extraction (dashboard/utils/resume.py)."""

import io
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "dashboard"))

from utils.resume import MAX_BYTES, ResumeReadError, extract_resume_text  # noqa: E402


def test_txt_is_read():
    assert extract_resume_text("cv.txt", b"Python and SQL") == "Python and SQL"


def test_docx_text_and_tables_are_read():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Skilled in Docker")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Kubernetes"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_resume_text("cv.docx", buf.getvalue())
    assert "Docker" in text and "Kubernetes" in text


def test_pdf_text_is_read():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "Experienced with Python and Tableau")
    text = extract_resume_text("cv.pdf", bytes(pdf.output()))
    assert "Python" in text and "Tableau" in text


def test_unsupported_type_is_rejected():
    with pytest.raises(ResumeReadError):
        extract_resume_text("cv.exe", b"data")


def test_oversize_file_is_rejected():
    with pytest.raises(ResumeReadError):
        extract_resume_text("cv.txt", b"a" * (MAX_BYTES + 1))


def test_corrupt_pdf_gives_friendly_error():
    with pytest.raises(ResumeReadError):
        extract_resume_text("cv.pdf", b"not really a pdf")


def test_empty_file_is_rejected():
    with pytest.raises(ResumeReadError):
        extract_resume_text("cv.txt", b"   ")
