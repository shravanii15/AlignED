"""Tests for generated PDF reports."""

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "dashboard"))

from services.reports_pdf import build_match_pdf_report  # noqa: E402


def test_match_pdf_is_a_valid_pdf():
    data = build_match_pdf_report(["Python", "Git"], [("Docker", 0.10), ("Kubernetes", 0.11)], 4)
    assert data.startswith(b"%PDF-") and len(data) > 1000


def test_match_pdf_handles_nothing_missing_and_unicode():
    data = build_match_pdf_report(["Python"], [], 1)
    assert data.startswith(b"%PDF-")
    data = build_match_pdf_report([], [("Café ☃ skill", 0.05)], 1)
    assert data.startswith(b"%PDF-")


def _sample_recs():
    import pandas as pd

    return pd.DataFrame([
        {"canonical_name": "Python", "market_demand_rate": 0.40, "program_coverage_rate": 0.0, "gap_value": 0.40,
         "trend_label": "no clear trend", "q_value": 0.0001, "p_value": 0.0001, "priority_tier": "high",
         "rationale": "Python appears in 40% of postings, but no course description names it."},
        {"canonical_name": "Docker", "market_demand_rate": 0.10, "program_coverage_rate": 0.004, "gap_value": 0.096,
         "trend_label": "rising", "q_value": 0.02, "p_value": 0.01, "priority_tier": "medium",
         "rationale": "Docker appears in 10% of postings."},
    ])


def test_skill_gap_pdf_builds_with_and_without_strengths():
    from services.reports_pdf import build_pdf_report

    recs = _sample_recs()
    assert build_pdf_report("Uni", "MS CS", 100, recs, "all jobs", 1660, 2).startswith(b"%PDF-")
    assert build_pdf_report("Uni", "MS CS", 100, recs, "all jobs", 1660, 2, strengths=[("C", 3)]).startswith(b"%PDF-")
    assert build_pdf_report("Uni", "MS CS", 100, recs.iloc[0:0], "all jobs", 1660, 0).startswith(b"%PDF-")


def test_excel_report_is_plain_language_and_keeps_numbers_numeric():
    import io

    from openpyxl import load_workbook
    from services.reports_excel import build_excel_report

    wb = load_workbook(io.BytesIO(build_excel_report("Uni", "MS CS", 100, _sample_recs())))
    assert wb.sheetnames == ["Overview", "Confirmed gaps", "How to read this"]
    ws = wb["Confirmed gaps"]
    headers = [c.value for c in ws[6]]
    assert "How common in jobs" in headers and "q-value" not in " ".join(str(h) for h in headers)
    assert ws["C7"].value == 0.40 and isinstance(ws["D7"].value, (int, float))


def _sample_picture():
    import pandas as pd

    return pd.DataFrame([
        {"skill_id": 1, "canonical_name": "Python", "demand": 0.40, "coverage": 0.0, "gap": 0.40, "status": "Not named"},
        {"skill_id": 2, "canonical_name": "SQL", "demand": 0.30, "coverage": 0.02, "gap": 0.28, "status": "Rarely named"},
        {"skill_id": 3, "canonical_name": "Git", "demand": 0.20, "coverage": 0.10, "gap": 0.10, "status": "Named"},
    ])


def test_reports_are_useful_even_with_zero_confirmed_gaps():
    import io

    from openpyxl import load_workbook
    from services.reports_excel import build_excel_report
    from services.reports_pdf import build_pdf_report

    empty = _sample_recs().iloc[0:0]
    wb = load_workbook(io.BytesIO(build_excel_report("Uni", "MS CS", 100, empty, picture=_sample_picture(), strengths=[("Git", 10)])))
    assert wb.sheetnames == ["Overview", "Top skills employers want", "Confirmed gaps", "How to read this"]
    assert wb["Top skills employers want"]["F5"].value == "Not named"
    assert build_pdf_report("Uni", "MS CS", 100, empty, "all jobs", 1660, 0, picture=_sample_picture()).startswith(b"%PDF-")


def test_market_picture_statuses():
    from utils.formatting import NAMED_WELL_AT, status_for

    assert status_for(0) == "Not named"
    assert status_for(NAMED_WELL_AT / 2) == "Rarely named"
    assert status_for(NAMED_WELL_AT) == "Named"
